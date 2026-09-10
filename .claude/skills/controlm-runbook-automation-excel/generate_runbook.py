"""DOC12 — fill the 2-tab Excel Application Run Book FROM THE GRAPH.

The sibling of ``generate_template.py``, and a different job. That one renders
``template-spec.yaml`` to an EMPTY workbook carrying the spec's ``example``
values. This one fills the real ones for a named Control-M folder, and marks
every cell it did not fill with WHY.

Run from the repo root:

    poetry run python .claude/skills/controlm-runbook-automation-excel/generate_runbook.py \
        --folder PRARAG-HLDM-70002-PEX-RFND-DLY --out <path>.xlsx

QUERIES ARE SHARED WITH THE -SDLC SIBLING, NEVER RE-WRITTEN. The L23 clause is
explicit that two skills disagreeing about one runbook fact is the failure to
prevent, so this module imports that generator's loaders and reads its
``FolderFacts`` rather than issuing Cypher of its own. There is exactly one
place in the repo that knows how to ask the graph what a folder is, and it is
not here. That import is by path because the skill is not on any package path —
the same mechanism ``tests/unit/test_sdlc_runbook_generator.py`` uses.

THREE OUTCOMES PER CELL, NOT TWO. The template declares a ``source:`` per field
— ``graph``, ``graph-partial``, ``manual`` — and the obvious reading is "fill
the graph ones, tint the rest". That reading is wrong in a way that matters:

* ``manual`` / ``graph-partial`` are DECLARED as needing a person. They tint,
  and say "needs capture". This is the source workbook's own yellow convention.
* a ``graph`` field this generator CAN fill gets its value.
* a ``graph`` field it CANNOT fill is the third state, and the one worth having.
  The spec claims the graph holds it; the graph, as DryDocs ingests it today,
  does not. Leaving that cell blank would be indistinguishable from "the folder
  genuinely has no value here", which is the silent-default disease this repo
  keeps paying for. It gets its own marker and its own count.

That third count is the honest measure of the spec's optimism, and it is
reported per tab at the end of every run. Where the reason is KNOWN it is named
from config rather than guessed: ``Avg Run time`` is empty because
``controlm@[db].psgmgr.cm_avg_run`` is confirmed and NOT WIRED on this side — a
wiring hold, in the vocabulary ``drydocs docs-coverage --section classes``
reports — and the cell says so instead of shrugging.

MECHANISM ONLY. This file ships no folder names, no ids and no values; a FILLED
workbook is Internal and belongs in ``internal/`` or ``internal-local/``, the
rule the -excel SKILL.md already carries.
"""

from __future__ import annotations

import argparse
import importlib.util
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
SPEC = HERE / "template-spec.yaml"
REPO = HERE.parents[2]
SDLC_GENERATOR = (
    REPO / ".claude" / "skills" / "controlm-runbook-automation-SDLC" / "generate_runbook.py"
)
DESCRIPTORS = REPO / "config" / "source-descriptors.yaml"

HEADER_FILL = PatternFill("solid", fgColor="FFFF00")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")
PARTIAL_FILL = PatternFill("solid", fgColor="FCE4B0")
#: The third state gets its OWN colour, because it is not the same request as a
#: yellow cell: yellow asks a person for a value, this asks the project for a feed.
UNREACHED_FILL = PatternFill("solid", fgColor="E6D0F5")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

NEEDS_CAPTURE = "needs capture"
#: Deliberately not the words "needs capture": nobody can capture this one by
#: hand into a graph that has no field for it.
NOT_REACHED = "declared graph-derivable; no feed reaches it here"

FILLED = "filled"
PARTIAL = "partial"
MANUAL = "manual"
UNREACHED = "unreached"


@dataclass
class Tally:
    """Per-tab counts, the run's own statement of what it managed.

    ``partial`` and ``manual`` are kept APART rather than summed into one
    "needs capture" number, because the template distinguishes them and the two
    are different asks: a graph-partial field has a named seam that needs
    enrichment, a manual one has no system of record at all. ``unreached`` is
    this generator's own addition and the number worth reading — see the module
    docstring's third outcome.
    """

    filled: int = 0
    partial: int = 0
    manual: int = 0
    unreached: int = 0

    def add(self, outcome: str) -> None:
        setattr(self, outcome, getattr(self, outcome) + 1)

    @property
    def total(self) -> int:
        return self.filled + self.partial + self.manual + self.unreached

    @property
    def declared_graph(self) -> int:
        """Fields the template says the graph holds: filled plus unreached."""
        return self.filled + self.unreached

    def render(self, name: str) -> str:
        return (
            f"{name}: {self.filled} filled, {self.partial} partial, "
            f"{self.manual} manual, {self.unreached} declared-graph-but-unreached "
            f"({self.filled}/{self.declared_graph} of the graph-declared fields "
            f"reached), {self.total} fields"
        )


@dataclass
class Filled:
    """One resolved cell: its value and which of the three outcomes produced it."""

    value: str
    outcome: str
    fill: PatternFill | None = None


@dataclass
class RunbookRun:
    """Everything one generation produced, for the caller and for the tests."""

    folder: str
    venue: str
    tallies: dict[str, Tally] = field(default_factory=dict)
    path: Path | None = None

    def render(self) -> str:
        lines = [f"folder {self.folder} (venue: {self.venue})"]
        lines += [f"  {t.render(name)}" for name, t in self.tallies.items()]
        if self.path:
            lines.append(f"  wrote {self.path}")
        return "\n".join(lines)


def _load_sdlc_generator():
    """The -SDLC sibling's loaders, imported by path.

    Shared rather than re-written: the L23 acceptance names two skills
    disagreeing about one runbook fact as the failure this prevents, so the
    Cypher lives there and only there.
    """
    spec = importlib.util.spec_from_file_location("_doc12_sdlc_facts", SDLC_GENERATOR)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _wired_reason(source_id: str) -> str:
    """The declared reason a dataset is not wired, read from config.

    Read, never restated, for the reason LOAD13's finish line reads the same
    block: the day the dataset is wired, this note stops being produced instead
    of persisting as a stale copy.
    """
    block = yaml.safe_load(DESCRIPTORS.read_text(encoding="utf-8")).get("wired") or {}
    entry = block.get(source_id)
    if isinstance(entry, dict) and entry.get("value") is False:
        return str(entry.get("reason") or "")
    return ""


#: A ``graph`` field whose emptiness has a KNOWN, declared cause. The value is
#: the registry dataset id that would carry it; the reason is read from config at
#: run time. Anything not here gets the generic unreached marker — a guess about
#: why would be worse than saying only that nothing reaches it.
KNOWN_HOLDS: dict[str, str] = {
    "Avg Run time": "controlm@[db].psgmgr.cm_avg_run",
}


def _technical_values(facts, module) -> dict[str, str]:
    """Technical_Details fields this generator can answer from FolderFacts.

    Every entry here is a fact the shared loaders already carry. A field absent
    from this map is NOT silently blank — it lands in the third state and is
    counted, which is what makes the gap visible instead of tidy.
    """
    app = facts.application or {}
    values = {
        "Control-M Folder name": facts.folder_name,
        "SEAL ID": str(app.get("app_id") or ""),
        "SEAL Application Name": str(app.get("name") or ""),
        "SOR Seal": str(app.get("app_id") or ""),
    }
    upstream = sorted(
        {
            folder
            for job in facts.jobs
            for wait in facts.conditions_for(job.get("job_id", ""))
            for folder, _job in facts.emitters.get(wait["condition_name"], [])
            if folder and folder != facts.folder_name
        }
    )
    if upstream:
        values["Dependency folder info in Control-M"] = ", ".join(upstream)
    scripts = facts.shell_scripts()
    if scripts:
        # `target` is the parser's answer, which is why one wrapper called by
        # four jobs is one entry here rather than four look-alikes.
        values["SCRIPT / Jar files location"] = ", ".join(sorted({s.target for s in scripts}))
    params = sorted({fact.config_path for fact in facts.commands() if fact.config_path})
    if params:
        values["Parameter file path"] = ", ".join(params)
    return {k: v for k, v in values.items() if v}


def _job_values(facts, job, module) -> dict[str, str]:
    """Control M Job details columns answerable per job from FolderFacts."""
    values = {
        "Folder Name": facts.folder_name,
        "Job Name/ SQL Agent Job": job.get("job_name", ""),
        "Description": job.get("description", ""),
        "Control-M Command Line (other than FW jobs)": job.get("cmd_line", ""),
        "CYCLIC (Y/N)": "Y" if (job.get("cyclic") or "").upper() in {"Y", "1", "TRUE"} else "N",
    }
    return {k: v for k, v in values.items() if v}


def _resolve(name: str, source: str, supplied: dict[str, str]) -> Filled:
    """The three-outcome rule, in one place so both tabs cannot diverge."""
    if source == "manual":
        return Filled(NEEDS_CAPTURE, MANUAL, MANUAL_FILL)
    if source == "graph-partial":
        return Filled(NEEDS_CAPTURE, PARTIAL, PARTIAL_FILL)
    value = supplied.get(name, "")
    if value:
        return Filled(value, FILLED, None)
    held = KNOWN_HOLDS.get(name)
    reason = _wired_reason(held) if held else ""
    note = f"{NOT_REACHED} - {held}: {reason}" if reason else NOT_REACHED
    return Filled(note, UNREACHED, UNREACHED_FILL)


def _write_technical(ws, tab: dict, facts, module, tally: Tally) -> None:
    supplied = _technical_values(facts, module)
    ws.append(tab["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
    ws.append(["", "", tab["legend"], ""])
    for row in tab["rows"]:
        resolved = _resolve(row["info"], row["source"], supplied)
        tally.add(resolved.outcome)
        ws.append([row["info"], resolved.value, row.get("comment", ""), ""])
        for cell in ws[ws.max_row]:
            cell.border = BORDER
            cell.alignment = WRAP
            if resolved.fill:
                cell.fill = resolved.fill
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 10
    ws.freeze_panes = "A2"


def _write_jobs(ws, tab: dict, facts, module, tally: Tally) -> None:
    cols = tab["columns"]
    ws.append([c["name"] for c in cols])
    for i, (cell, col) in enumerate(zip(ws[1], cols, strict=True), start=1):
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(i)].width = max(18, min(46, len(col["name"]) + 4))
    # The tally counts COLUMNS once, not once per job: the question the report
    # answers is which of the template's fields this generator can fill, and
    # multiplying that by the job count would inflate it by an accident of the
    # folder's size.
    counted = False
    for job in facts.jobs:
        supplied = _job_values(facts, job, module)
        ws.append([_resolve(c["name"], c["source"], supplied).value for c in cols])
        for cell, col in zip(ws[ws.max_row], cols, strict=True):
            resolved = _resolve(col["name"], col["source"], supplied)
            cell.border = BORDER
            cell.alignment = WRAP
            if resolved.fill:
                cell.fill = resolved.fill
            if not counted:
                tally.add(resolved.outcome)
        counted = True
    if not facts.jobs:
        # A folder with no jobs still declares the shape. Count the columns off
        # an empty supply so the tab reports its fields rather than reporting
        # nothing at all.
        for col in cols:
            tally.add(_resolve(col["name"], col["source"], {}).outcome)
    ws.freeze_panes = "C2"


def generate(
    folder: str,
    *,
    facts=None,
    venue: str = "",
    out: Path | None = None,
) -> RunbookRun:
    """Fill the workbook for one folder. Returns what the run produced.

    ``facts`` is INJECTED rather than fetched when the caller already has them —
    which is how the graph path works: the caller builds a client, calls the
    shared ``load_from_graph``, and hands the result here. This module never
    constructs a Neo4j client, so there is no second place that decides how the
    graph is reached. With no facts given it falls back to the bundled samples,
    which is what makes the generator demonstrable in any clone.
    """
    module = _load_sdlc_generator()
    if facts is None:
        facts = module.load_from_samples(folder)
        venue = venue or "bundled sample CSVs"
    spec = yaml.safe_load(SPEC.read_text(encoding="utf-8"))
    run = RunbookRun(folder=folder, venue=venue or getattr(facts, "venue", "") or "unstated")
    wb = Workbook()
    wb.remove(wb.active)
    for tab in spec["tabs"]:
        ws = wb.create_sheet(tab["name"])
        tally = Tally()
        if tab["layout"] == "rows":
            _write_technical(ws, tab, facts, module, tally)
        else:
            _write_jobs(ws, tab, facts, module, tally)
        run.tallies[tab["name"]] = tally
    if out is not None:
        wb.save(out)
        run.path = out
    return run


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--folder", required=True, help="the Control-M folder (sched_table) name")
    ap.add_argument("--out", type=Path, default=None, help="write the workbook here")
    args = ap.parse_args(argv)
    run = generate(args.folder, out=args.out)
    print(run.render())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
